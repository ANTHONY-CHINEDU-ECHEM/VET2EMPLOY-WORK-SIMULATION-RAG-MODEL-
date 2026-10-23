"""Workbook parsing, file metadata, document parsing and chunking."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import docx
import pytest
from openpyxl import Workbook

from tests.conftest import ROOT
from v2e_rag.domain.catalog import load_catalog
from v2e_rag.domain.models import SourceDocument
from v2e_rag.domain.taxonomy import Facet, JobLevel, SourceType, normalise_job_level
from v2e_rag.ingestion.chunking import (
    DocumentChunker,
    RoleChunker,
    activity_number_from_path,
    pack_parts,
)
from v2e_rag.ingestion.document_parser import (
    Block,
    assemble_paragraphs,
    join_cell_lines,
    linearise_table,
    parse_docx,
)
from v2e_rag.ingestion.excel_loader import (
    SheetKind,
    WorkbookLoader,
    canonical_header,
    classify_sheet,
)
from v2e_rag.ingestion.filename_metadata import (
    infer_file_metadata,
    parse_collection_date,
    strip_upload_prefix,
)
from v2e_rag.ingestion.pipeline import document_id_for

CATALOG = load_catalog(str(ROOT / "config"))


# File names -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "families", "sectors", "collected"),
    [
        (
            "Cybersecurity_Analyst_Construction_Oct_22nd_2026.xlsx",
            ("cybersecurity",),
            ("construction",),
            date(2026, 10, 22),
        ),
        (
            "Data_Analyst__Consulting_21st_Oct_2026.xlsx",
            ("data_analyst",),
            ("consulting",),
            date(2026, 10, 21),
        ),
        (
            "Healthcare_Data_Analytics_2026-09-29.xlsx",
            ("data_analyst",),
            ("healthcare",),
            date(2026, 9, 29),
        ),
        (
            "Healthcare_Data_Engineering_1st_of_October_2026.xlsx",
            ("data_engineer",),
            ("healthcare",),
            date(2026, 10, 1),
        ),
        (
            "IT_Software_Developer_Oct_19th_2026.xlsx",
            ("software_developer",),
            ("information_technology",),
            date(2026, 10, 19),
        ),
        (
            "Cybersecurity_Education_sector__15th_October_2026.xlsx",
            ("cybersecurity",),
            ("education",),
            date(2026, 10, 15),
        ),
        (
            "Software_developer_September_24th_2026.xlsx",
            ("software_developer",),
            ("general",),
            date(2026, 9, 24),
        ),
    ],
)
def test_infer_file_metadata(name, families, sectors, collected):
    meta = infer_file_metadata("3ef82fe0-1792788177881_" + name)
    assert meta.clean_name == name
    assert meta.role_families == families
    assert meta.sectors == sectors
    assert meta.collected_on == collected


def test_upload_prefix_and_document_id():
    assert strip_upload_prefix("015cca85-1792788177881_File.xlsx") == "File.xlsx"
    assert (
        document_id_for("015cca85-1792788177881_Data_Analyst__Consulting.xlsx")
        == "doc_data_analyst_consulting"
    )
    assert parse_collection_date("no_date_here") is None


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Mid-level", JobLevel.MID_LEVEL),
        ("Mid-Level", JobLevel.MID_LEVEL),
        ("Beginner ", JobLevel.BEGINNER),
        ("Advanced ", JobLevel.ADVANCED),
        ("Expert", JobLevel.EXPERT),
        ("Wizard", None),
        (None, None),
    ],
)
def test_job_level_normalisation(raw, expected):
    assert normalise_job_level(raw) == expected


# Workbooks --------------------------------------------------------------------------------


def _workbook(path: Path) -> Path:
    wb = Workbook()
    master = wb.active
    master.title = "Master dataset"
    master.append(
        [
            "Company ID",
            "Company",
            "Niche",
            "Location",
            "Job Title",
            "Job Level",
            "Skills Required",
            "Role / Day-to-Day Responsibilities",
            "Tech Stack",
            "Team Collaboration & Culture",
            "No of Reviews",
            "Overall Company Rating",
            "Culture Rating",
            "Source URLs",
        ]
    )
    master.append(
        [
            "ACM-001",
            "Acme Ltd",
            "Smart Homes",
            "Leeds",
            "Data Analyst",
            "Mid-Level",
            "SQL; Excel",
            "Build dashboards; Profile data",
            "Power BI; SQL",
            "Works with finance",
            "1,234",
            "4.1",
            "N/A",
            "https://example.com/job/1",
        ]
    )
    master.append(
        [
            "NUF-002",
            "Nuffield",
            "Health",
            "York",
            "Data Protection Analyst",
            "Advanced",
            "GDPR",
            "Review DPIAs",
            "OneTrust",
            "Works with legal",
            None,
            None,
            None,
            "https://example.com/job/2",
        ]
    )
    topic = wb.create_sheet("Tech stack")
    topic.append(
        [
            "Company ID",
            "Company Name",
            "Job Title",
            "Job level",
            "Frameworks",
            "Tech Stack",
            "Source URL",
        ]
    )
    topic.append(
        [
            "ACM-001",
            "Acme Ltd",
            "Data Analyst",
            "Mid-level",
            "DAMA DMBOK",
            "Ignored duplicate",
            "https://example.com/job/1b",
        ]
    )
    topic.append(
        ["NUF-0022", "Nuffield", "Data Protection Analyst", "Advanced", "ISO 27701", None, None]
    )
    culture = wb.create_sheet("Work culture ")
    culture.append(
        [
            "Company ID",
            "Company Name",
            "Job Title",
            "Job level",
            "No of Reviews",
            "Overall Company Rating",
        ]
    )
    culture.append(["ZED-001", "Zed Group", "Analyst", "Beginner", "10", "3.5"])
    method = wb.create_sheet("Rating methodology")
    method.append(["Indeed Rating Methodology"])
    method.append(["Core Principle", "Weighted mean of star ratings"])
    wb.save(path)
    return path


def test_workbook_loader_merges_sheets(tmp_path):
    path = _workbook(tmp_path / "Data_Analyst_Construction_Oct_22nd_2026.xlsx")
    result = WorkbookLoader(path, "doc_x", infer_file_metadata(path.name)).load()
    by_title = {r.job_title: r for r in result.records}

    acme = by_title["Data Analyst"]
    assert acme.job_level == JobLevel.MID_LEVEL
    assert acme.fields["review_count"] == 1234
    assert acme.fields["rating_overall"] == 4.1
    assert "rating_culture" not in acme.fields  # N/A is dropped, never guessed
    assert acme.fields["tech_stack"] == "Power BI; SQL"  # master wins over topic sheet
    assert acme.fields["frameworks"] == "DAMA DMBOK"  # topic sheet adds missing fields
    assert acme.source_urls == ["https://example.com/job/1", "https://example.com/job/1b"]

    # A mistyped Company ID on a topic sheet still reattaches by employer and title.
    assert by_title["Data Protection Analyst"].fields["frameworks"] == "ISO 27701"

    culture_only = [r for r in result.records if r.fields.get("record_kind") == "company_culture"]
    assert len(culture_only) == 1 and culture_only[0].company == "Zed Group"
    assert result.references and result.references[0].kind == SheetKind.METHODOLOGY
    assert result.sheet_kinds["Work culture "] == "ratings"


def test_header_aliases_and_classification():
    assert canonical_header("Team Collaboration & Culture") == "team_collaboration"
    assert canonical_header(" Source URLs ") == "source_url"
    assert canonical_header("Role / Day\u2011to\u2011Day Responsibilities") == "responsibilities"
    assert classify_sheet("MasterDataset", set()) == SheetKind.MASTER
    assert classify_sheet("Company Reviews", {"employee_review", "job_title"}) == SheetKind.REVIEWS
    assert classify_sheet("Employee Perspectives", set()) == SheetKind.PERSPECTIVES


# Documents --------------------------------------------------------------------------------


def test_table_cell_and_paragraph_assembly():
    assert (
        join_cell_lines("PDF,\nV2E_DA_[Surname]_D\n1_AnalysisPlan_v1.pdf")
        == "PDF, V2E_DA_[Surname]_D1_AnalysisPlan_v1.pdf"
    )
    assert join_cell_lines("Questions, data,\nschedule") == "Questions, data, schedule"
    lines = linearise_table([["Ref", "Due"], ["D1", "Mon\n17:00"]])
    assert lines == ["Ref: D1 | Due: Mon 17:00"]
    text = assemble_paragraphs(["You must:", "\u2022 do one", "thing well", "1. then this"])
    assert text.split("\n") == ["You must:", "\u2022 do one thing well", "1. then this"]


def test_docx_parsing_uses_heading_styles(tmp_path):
    document = docx.Document()
    document.add_heading("Part 1 The Problem", level=1)
    document.add_paragraph("The experience paradox affects graduates.")
    document.add_heading("What Employers Want", level=2)
    document.add_paragraph("Evidence of professional behaviour.", style="List Paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Role", "Builds"
    table.cell(1, 0).text, table.cell(1, 1).text = "Developer", "The platform"
    path = tmp_path / "report.docx"
    document.save(path)
    blocks = parse_docx(path)
    assert [b.kind for b in blocks] == ["heading", "text", "heading", "text", "table"]
    assert blocks[2].level == 2 and blocks[3].text.startswith("\u2022 ")


def test_section_chunker_tags_tracks_facets_and_activities():
    blocks = [
        Block("heading", 5, "3. How Your Week Runs", 2),
        Block("text", 5, "Escalate after two hours."),
        Block("heading", 17, "Data Analyst", 1),
        Block("heading", 18, "7. Track 1: Data Analyst", 2),
        Block("heading", 22, "7.6 Sequential activities", 4),
        Block("heading", 24, "Activity 3: Data profiling and quality audit", 6),
        Block("text", 24, "Record each issue in the Data Quality Log."),
        Block("heading", 31, "7.7 Required deliverables", 4),
        Block("table", 31, rows=[["Ref", "Deliverable"], ["D2", "Data Quality Log"]]),
        Block("heading", 80, "11. Appendices", 2),
        Block("text", 80, "Appendix B daily update template."),
    ]
    document = SourceDocument(
        document_id="brief",
        source_file="brief.pdf",
        title="Brief",
        source_type=SourceType.PROGRAMME_BRIEF,
        content_hash="x",
    )
    chunks = DocumentChunker(CATALOG, 400, 40).chunk(document, blocks)
    by_text = {c.body: c for c in chunks}
    norms = by_text["Escalate after two hours."]
    assert norms.track is None and norms.facet == Facet.PROGRAMME_RULES
    activity = by_text["Record each issue in the Data Quality Log."]
    assert activity.track == "DA" and activity.facet == Facet.TRACK_ACTIVITY
    assert activity.metadata["activity_number"] == 3
    assert "Track DA: Data Analyst" in activity.header
    deliverable = by_text["Ref: D2 | Deliverable: Data Quality Log"]
    assert deliverable.facet == Facet.TRACK_DELIVERABLES and deliverable.page_start == 31
    appendix = by_text["Appendix B daily update template."]
    assert appendix.track is None and appendix.facet == Facet.TEMPLATES


def test_pack_parts_respects_budget_and_overlap():
    parts = [f"Sentence number {i} has several words in it." for i in range(60)]
    windows = pack_parts(parts, max_tokens=60, overlap_tokens=12)
    assert len(windows) > 3
    assert all(len(w.split()) <= 60 for w in windows)
    assert windows[1].split("\n")[0] in windows[0]  # overlap carries context forward
    assert activity_number_from_path("7.6 Sequential activities > Activity 12: Handover") == 12


def test_role_chunker_builds_facets(tmp_path):
    path = _workbook(tmp_path / "Data_Analyst_Construction_Oct_22nd_2026.xlsx")
    result = WorkbookLoader(path, "doc_x", infer_file_metadata(path.name)).load()
    document = SourceDocument(
        document_id="doc_x",
        source_file=path.name,
        title="t",
        source_type=SourceType.JOB_MARKET,
        content_hash="x",
    )
    chunks = RoleChunker(400).chunk(document, result.records[0])
    facets = {c.facet for c in chunks}
    assert {
        Facet.RESPONSIBILITIES,
        Facet.SKILLS_EXPERIENCE,
        Facet.TECH_STACK,
        Facet.TEAM_COLLABORATION,
        Facet.CULTURE_RATINGS,
    } <= facets
    ratings = next(c for c in chunks if c.facet == Facet.CULTURE_RATINGS)
    assert (
        ratings.source_type == SourceType.EMPLOYER_REVIEW
        and "1234 employee reviews" in ratings.body
    )
    assert all("Level: Mid Level" in c.header and c.company == "Acme Ltd" for c in chunks)
